"""
Mantri Salesforce lead status report.

The raw Salesforce Excel export is UPLOADED via the InsightDesk UI and stored
in the database (mis_salesforce_uploads) — no local file access, so this works
on ephemeral hosting (Railway). The report parses the latest uploaded file,
maps Sub Source to Meta/Google, and returns lead-status counts/percentages.
"""
import io
import logging
from datetime import datetime, timedelta, date as date_type
from typing import Any, Dict, List, Optional

import openpyxl
from fastapi import APIRouter, Depends, HTTPException, UploadFile, File
from sqlalchemy.orm import Session

from backend.db.database import get_db
from backend.db.models import MisProject, MisSalesforceUpload, User
from backend.routes.auth import get_current_user_required
from backend.services.activity_log import log_activity

logger = logging.getLogger("AdOptima")

router = APIRouter(prefix="/api/mis/mantri/salesforce", tags=["salesforce_mantri"])

MAX_UPLOAD_BYTES = 20 * 1024 * 1024  # 20 MB

# Sub Source -> platform mapping
SUB_SOURCE_MAP = {
    "Facebook": "Meta",
    "Facebook Lead Page": "Meta",
    "Landing Page": "Meta",
    "google_paid": "Google",
}
DEFAULT_PLATFORM = "Meta"

# Lead statuses in the order shown in the reference report
LEAD_STATUS_ORDER = [
    "Booked",
    "Closed Lost",
    "Interested",
    "New",
    "Not Interested",
    "Post Site Visit Follow-Up",
    "Pre Sales Follow Up",
    "Qualified",
    "Sales Follow up",
    "Site Visit Schedule",
    "SV Completed",
]

# Fixed report window: 16-04-2026 to yesterday (today always excluded)
REPORT_START = datetime.strptime("16-04-2026", "%d-%m-%Y").date()


def _map_platform(sub_source: Any) -> str:
    if sub_source is None:
        return DEFAULT_PLATFORM
    val = str(sub_source).strip()
    if not val or val.lower() == "nan":
        return DEFAULT_PLATFORM
    return SUB_SOURCE_MAP.get(val, DEFAULT_PLATFORM)


def _parse_create_date(value: Any) -> Optional[date_type]:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date_type):
        return value
    if hasattr(value, "date") and callable(value.date):
        try:
            return value.date()
        except Exception:
            pass
    val_str = str(value).strip()
    if not val_str or val_str.lower() == "nan":
        return None
    for fmt in ("%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d", "%m/%d/%Y"):
        try:
            return datetime.strptime(val_str.split()[0], fmt).date()
        except Exception:
            continue
    return None


def _parse_xlsx_bytes(content: bytes, start_date: date_type, end_date: date_type) -> List[Dict[str, Any]]:
    """Parse the Salesforce export bytes and return filtered lead rows."""
    wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    try:
        ws = wb.active
        rows_iter = ws.iter_rows(values_only=True)
        # Skip to header row (Salesforce exports have metadata rows above; header at row 10)
        for _ in range(9):
            next(rows_iter, None)
        header_row = next(rows_iter, None)
        if not header_row:
            raise ValueError("Could not find header row in the uploaded file (expected row 10)")
        headers = [str(c).strip() if c is not None else "" for c in header_row]
        col_indices = {}
        for idx, h in enumerate(headers):
            if h in ("Create Date", "Sub Source", "Lead status"):
                col_indices[h] = idx

        missing = [h for h in ("Create Date", "Sub Source", "Lead status") if h not in col_indices]
        if missing:
            raise ValueError(f"Uploaded file is missing required column(s): {', '.join(missing)}")

        create_date_idx = col_indices["Create Date"]
        sub_source_idx = col_indices["Sub Source"]
        lead_status_idx = col_indices["Lead status"]

        leads = []
        for row in rows_iter:
            if not row:
                continue
            raw_date = row[create_date_idx] if create_date_idx < len(row) else None
            parsed_d = _parse_create_date(raw_date)
            if not parsed_d or parsed_d < start_date or parsed_d > end_date:
                continue
            raw_status = row[lead_status_idx] if lead_status_idx < len(row) else None
            if raw_status is None or not str(raw_status).strip() or str(raw_status).strip().lower() == "nan":
                continue
            raw_sub_source = row[sub_source_idx] if sub_source_idx < len(row) else None
            leads.append({
                "create_date": parsed_d,
                "lead_status": str(raw_status).strip(),
                "platform": _map_platform(raw_sub_source),
            })
        return leads
    finally:
        wb.close()


def _get_mantri_project(db: Session) -> Optional[MisProject]:
    return (
        db.query(MisProject)
        .filter(MisProject.name.ilike("%serenity%") | MisProject.name.ilike("%mantri%"))
        .first()
    )


def _compute_report(leads: List[Dict[str, Any]], start_date: date_type, end_date: date_type) -> Dict[str, Any]:
    status_counts: Dict[str, Dict[str, int]] = {status: {"Meta": 0, "Google": 0} for status in LEAD_STATUS_ORDER}

    for item in leads:
        status = item["lead_status"]
        platform = item["platform"]
        if status in status_counts and platform in status_counts[status]:
            status_counts[status][platform] += 1

    meta_total = sum(status_counts[s]["Meta"] for s in LEAD_STATUS_ORDER)
    google_total = sum(status_counts[s]["Google"] for s in LEAD_STATUS_ORDER)
    overall_total = meta_total + google_total

    rows = []
    for status in LEAD_STATUS_ORDER:
        meta_count = status_counts[status]["Meta"]
        google_count = status_counts[status]["Google"]
        overall_count = meta_count + google_count
        rows.append({
            "lead_status": status,
            "meta_count": meta_count,
            "meta_pct": round((meta_count / meta_total * 100), 1) if meta_total else 0.0,
            "google_count": google_count,
            "google_pct": round((google_count / google_total * 100), 1) if google_total else 0.0,
            "overall_count": overall_count,
            "overall_pct": round((overall_count / overall_total * 100), 1) if overall_total else 0.0,
        })

    return {
        "title": "MANTRI - SALESFORCE LEAD STATUS",
        "start_date": start_date.isoformat(),
        "end_date": end_date.isoformat(),
        "rows": rows,
        "totals": {
            "meta_count": meta_total,
            "google_count": google_total,
            "overall_count": overall_total,
        },
    }


def _latest_upload(db: Session, project_id: int) -> Optional[MisSalesforceUpload]:
    return (
        db.query(MisSalesforceUpload)
        .filter(MisSalesforceUpload.project_id == project_id)
        .order_by(MisSalesforceUpload.uploaded_at.desc())
        .first()
    )


@router.post("/upload")
async def upload_salesforce_file(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user_required),
):
    """Upload the raw Salesforce Excel export (replaces the previous upload for reports)."""
    if not file.filename or not file.filename.lower().endswith((".xlsx", ".xls")):
        raise HTTPException(status_code=400, detail="Please upload a Salesforce Excel export (.xlsx)")

    content = await file.read()
    if len(content) > MAX_UPLOAD_BYTES:
        raise HTTPException(status_code=400, detail="File too large (max 20 MB)")
    if not content:
        raise HTTPException(status_code=400, detail="Uploaded file is empty")

    # Validate it parses as xlsx BEFORE storing
    try:
        openpyxl.load_workbook(io.BytesIO(content), read_only=True).close()
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Invalid Excel file: {e}")

    project = _get_mantri_project(db)
    if not project:
        raise HTTPException(status_code=404, detail="Mantri MIS project not found")

    upload = MisSalesforceUpload(
        project_id=project.id,
        filename=file.filename,
        file_size=len(content),
        content=content,
        uploaded_by=user.email,
    )
    db.add(upload)
    db.commit()
    db.refresh(upload)

    log_activity(
        module="InsightDesk",
        action="Mantri Salesforce File Uploaded",
        description=f"Uploaded Salesforce export '{file.filename}' ({len(content)} bytes)",
        user_id=user.id,
        user_name=user.full_name or user.email,
        account_id=project.client_id,
        entity_type="mis_project",
        entity_id=str(project.id),
        details={"filename": file.filename, "file_size": len(content)},
        db=db,
    )

    return {
        "status": "success",
        "upload": upload.to_dict(),
        "message": "File attached. Generating report...",
    }


@router.get("/upload/status")
def upload_status(
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user_required),
):
    """Return info about the latest uploaded Salesforce file (if any)."""
    project = _get_mantri_project(db)
    if not project:
        return {"has_upload": False}
    upload = _latest_upload(db, project.id)
    if not upload:
        return {"has_upload": False}
    return {"has_upload": True, "upload": upload.to_dict()}


@router.get("/report")
def salesforce_report(
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user_required),
):
    """Return Mantri Salesforce lead status report from the latest uploaded file.

    Date range is fixed: 16-04-2026 to yesterday (inclusive). Today is excluded.
    """
    project = _get_mantri_project(db)
    if not project:
        raise HTTPException(status_code=404, detail="Mantri MIS project not found")

    upload = _latest_upload(db, project.id)
    if not upload:
        raise HTTPException(
            status_code=404,
            detail="No Salesforce file attached yet. Click 'Attach Raw File' to upload the export from CRM.",
        )

    end_date = datetime.utcnow().date() - timedelta(days=1)
    try:
        leads = _parse_xlsx_bytes(upload.content, REPORT_START, end_date)
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))
    except Exception as e:
        logger.exception("Failed to parse uploaded Salesforce file")
        raise HTTPException(status_code=400, detail=f"Failed to read the uploaded file: {e}")

    if not leads:
        raise HTTPException(
            status_code=404,
            detail=f"No leads found in the report window ({REPORT_START.strftime('%d-%m-%Y')} to {end_date.strftime('%d-%m-%Y')}). Check the date filters in your CRM export.",
        )

    report = _compute_report(leads, REPORT_START, end_date)
    report["source_file"] = {"filename": upload.filename, "uploaded_at": upload.uploaded_at.isoformat() if upload.uploaded_at else None}

    log_activity(
        module="InsightDesk",
        action="Mantri Salesforce Report Generated",
        description=f"Generated Mantri Salesforce lead status report from '{upload.filename}' ({len(leads)} leads, {REPORT_START.isoformat()} to {end_date.isoformat()})",
        user_id=user.id,
        user_name=user.full_name or user.email,
        account_id=project.client_id,
        entity_type="mis_project",
        entity_id=str(project.id),
        details={"leads_count": len(leads), "filename": upload.filename},
        db=db,
    )

    return report