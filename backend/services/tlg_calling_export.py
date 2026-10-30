"""Native Excel export of the selected, authorized calling MIS snapshot."""
from datetime import datetime, timedelta, timezone
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

COLUMNS = [('Centre', 'centre', None), ('Leads', 'leads', None),
    ('Called', 'called', None), ('%', 'called', 'leads'),
    ('Not called', 'not_called', None), ('%', 'not_called', 'leads'),
    ('Contacted', 'contacted', None), ('%', 'contacted', 'called'),
    ('Not contactable', 'not_contactable', None), ('%', 'not_contactable', 'called'),
    ('Interested', 'interested', None), ('%', 'interested', 'contacted'),
    ('Not interested', 'not_interested', None), ('%', 'not_interested', 'contacted'),
    ('Junk', 'junk', None), ('%', 'junk', 'contacted'),
    ('Pending / Unclassified', 'pending', None), ('%', 'pending', 'contacted'),
    ('Scheduled', 'scheduled', None), ('%', 'scheduled', 'interested'),
    ('Not scheduled', 'not_scheduled', None), ('%', 'not_scheduled', 'interested'),
    ('Visited', 'visited', None), ('%', 'visited', 'scheduled'),
    ('Not visited', 'not_visited', None), ('%', 'not_visited', 'scheduled')]

def calling_analysis_workbook(report):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Calling MIS'
    ws.merge_cells('A1:Z1')
    ws['A1'] = 'The Little Gym — Centrewise Calling MIS'
    ws['A1'].font = Font(name='Arial', size=16, bold=True, color='0F172A')
    ws.merge_cells('A2:Z2')
    ws['A2'] = f"Selected dates: {report['start']} to {report['end']}"
    checked = datetime.fromisoformat(report['checked_at'].replace('Z', '+00:00')).astimezone(timezone(timedelta(hours=5, minutes=30)))
    ws.merge_cells('A3:Z3')
    ws['A3'] = f"Snapshot updated: {checked:%d %b %Y, %I:%M %p} IST" + (' — older report; refresh overdue' if report.get('stale') else '')
    groups = [(1, 10, 'Leads Outcome', 'DBEAFE'), (11, 18, 'Outcomes of contacted leads', 'EDE9FE'), (19, 26, 'Visit outcomes', 'D1FAE5')]
    for first, last, title, color in groups:
        ws.merge_cells(start_row=4, start_column=first, end_row=4, end_column=last)
        ws.cell(4, first, title)
        for row in ws.iter_rows(min_row=4, max_row=5, min_col=first, max_col=last):
            for cell in row:
                cell.fill = PatternFill('solid', fgColor=color)
                cell.font = Font(name='Arial', size=11, bold=True, color='0F172A')
                cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    for col, (label, _, _) in enumerate(COLUMNS, 1):
        ws.cell(5, col, label)
        ws.column_dimensions[get_column_letter(col)].width = 8 if label == '%' else 17
    ws.column_dimensions['A'].width = 19
    ws.column_dimensions['B'].width = 10
    ws.column_dimensions['Q'].width = 25
    border = Border(bottom=Side(style='thin', color='CBD5E1'), right=Side(style='thin', color='CBD5E1'))
    for row_num, data in enumerate([*report['rows'], report['total']], 6):
        is_total = row_num == 6 + len(report['rows'])
        for col, (_, key, denominator) in enumerate(COLUMNS, 1):
            value = (data[key] / data[denominator] if data[denominator] else 0) if denominator else data[key]
            cell = ws.cell(row_num, col, value)
            cell.font = Font(name='Arial', size=12 if is_total else 11, bold=is_total, color='0F172A')
            cell.fill = PatternFill('solid', fgColor='E2E8F0' if is_total else ('EFF6FF' if col <= 10 else 'F5F3FF' if col <= 18 else 'ECFDF5'))
            cell.number_format = '0.0%' if denominator else ('General' if col == 1 else '0')
            cell.alignment = Alignment(horizontal='left' if col == 1 else 'center', vertical='center')
            cell.border = border
        ws.row_dimensions[row_num].height = 25
    footer = 6 + len(report['rows'])
    notes = [
        'Percentages: Called / Not called ÷ Leads; Contacted / Not contactable ÷ Called; outcomes ÷ Contacted; Scheduled / Not scheduled ÷ Interested; Visited / Not visited ÷ Scheduled.',
        'Wrong / invalid numbers remain Not contactable. Scheduled includes completed visits. Test names/emails excluded.']
    if any(row.get('unknown_dates') for row in report['rows']):
        notes.append('Unparseable lead timestamps are included for continuity, as in the displayed report.')
    for row_num, note in enumerate(notes, footer + 2):
        ws.merge_cells(start_row=row_num, start_column=1, end_row=row_num, end_column=26)
        ws.cell(row_num, 1, note).font = Font(name='Arial', size=10, color='334155')
    ws.freeze_panes = 'C6'
    ws.auto_filter.ref = f'A5:Z{footer - 1}'
    ws.row_dimensions[1].height = 30
    ws.row_dimensions[4].height = 25
    ws.row_dimensions[5].height = 38
    # Centre labels are text, even if an account administrator uses formula-like names.
    for row in ws:
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = 's'
    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)
    return stream
