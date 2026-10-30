import unittest
from openpyxl import load_workbook
from backend.services.tlg_calling_analysis import METRICS
from backend.services.tlg_calling_export import calling_analysis_workbook

class ExportTests(unittest.TestCase):
    def report(self):
        first={'centre':'Jayanagar',**dict.fromkeys(METRICS,0),'leads':4,'called':3,'not_called':1,'contacted':2,'not_contactable':1,'interested':1,'pending':1,'scheduled':1,'visited':1}
        second={'centre':'Whitefield',**dict.fromkeys(METRICS,0),'leads':1,'not_called':1}
        total={'centre':'Total',**{key:first[key]+second[key] for key in METRICS}}
        return {'rows':[first,second],'total':total,'start':'2026-10-01','end':'2026-10-29','checked_at':'2026-10-30T06:42:00Z','stale':False}
    def test_native_xlsx_counts_percentages_and_weighted_total(self):
        wb=load_workbook(calling_analysis_workbook(self.report()),data_only=True)
        ws=wb['Calling MIS']
        self.assertEqual(ws['B6'].value,4)
        self.assertEqual(ws['D6'].value,.75)
        self.assertEqual(ws['D7'].value,0)
        self.assertEqual(ws['D8'].value,.6)
        self.assertEqual(ws['H6'].value,2/3)
        self.assertEqual(ws['X6'].value,1)
        self.assertEqual(ws['D6'].number_format,'0.0%')
        self.assertEqual(ws.freeze_panes,'C6')
        self.assertEqual(ws.auto_filter.ref,'A5:Z7')
        self.assertEqual(len(wb.sheetnames),1)
    def test_visible_period_verification_time_and_contrast(self):
        ws=load_workbook(calling_analysis_workbook(self.report())).active
        self.assertIn('2026-10-01 to 2026-10-29',ws['A2'].value)
        self.assertIn('12:12 PM',ws['A3'].value)
        self.assertEqual(ws['A5'].font.color.rgb,'000F172A')
        self.assertTrue(ws['A8'].font.bold)
        self.assertEqual(ws['A8'].font.color.rgb,'000F172A')
    def test_formula_like_centre_is_exported_as_text(self):
        data=self.report();data['rows'][0]['centre']='=1+1'
        ws=load_workbook(calling_analysis_workbook(data)).active
        self.assertEqual(ws['A6'].data_type,'s')
        self.assertEqual(ws['A6'].value,'=1+1')

if __name__=='__main__':unittest.main()
